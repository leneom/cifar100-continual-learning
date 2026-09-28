from __future__ import annotations

import csv
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "report"
TMP_DIR = ROOT / "tmp" / "report"
OUTPUT_DOCX = REPORT_DIR / "continual_learning_cifar100_report.docx"
CHART_PNG = TMP_DIR / "replay_multiseed.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "16324F"
MUTED = "5C6773"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
BORDER = "C9D2DC"
WHITE = "FFFFFF"
BLACK = "111827"

# Windows fonts, also reachable from WSL through the /mnt/c mount.
FONT_DIR = next(
    (path for path in (Path("C:/Windows/Fonts"), Path("/mnt/c/Windows/Fonts")) if path.exists()),
    Path("C:/Windows/Fonts"),
)


def set_run_font(run, size=None, bold=None, italic=None, color=BLACK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)


def style_table(table, numeric_columns=()):
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            if row_index == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if column_index in numeric_columns:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, size=9.2, bold=row_index == 0)


def add_table(doc, headers, rows, widths_dxa, numeric_columns=()):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = str(value)
    set_table_geometry(table, widths_dxa)
    style_table(table, numeric_columns=numeric_columns)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(2)
    after.paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run_node = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = "1"
    run_node.append(text_node)
    field.append(run_node)
    paragraph._p.append(field)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.4)  # Named override: four-page research brief body.
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(10.4)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.1

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.add_run("CONTINUAL LEARNING ON CIFAR-100")
    p.add_run("\tRESEARCH SUMMARY")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    for run in p.runs:
        set_run_font(run, size=8, bold=True, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run("Independent research project | August 2026")
    p.add_run("\t")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    set_run_font(p.runs[0], size=8.5, color=MUTED)
    add_page_number(p)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, size=10.4)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=10.4)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.4)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, size=10.4)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    set_run_font(run, size=10.4)
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_BLUE)
    p_pr.append(shading)
    label_run = p.add_run(f"{label}  ")
    set_run_font(label_run, size=10.4, bold=True, color=NAVY)
    body_run = p.add_run(text)
    set_run_font(body_run, size=10.4, color=BLACK)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def load_summary():
    with (ROOT / "runs" / "replay-multiseed" / "summary.csv").open(
        encoding="utf-8", newline=""
    ) as handle:
        return list(csv.DictReader(handle))


def font(size, bold=False):
    candidates = [
        Path(FONT_DIR, "arialbd.ttf" if bold else "arial.ttf"),
        Path(FONT_DIR, "calibrib.ttf" if bold else "calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def make_chart(rows):
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    pil = lambda value: value if value.startswith("#") else f"#{value}"
    image = Image.new("RGB", (1600, 680), pil(WHITE))
    draw = ImageDraw.Draw(image)
    draw.text((80, 42), "Replay capacity across three seeds", font=font(42, True), fill=pil(NAVY))
    draw.text(
        (80, 96),
        "Points show means; error bars show sample standard deviation (seeds 0, 1, and 42)",
        font=font(22),
        fill=pil(MUTED),
    )

    panels = [
        (80, 160, 690, 400, "Final average accuracy", "accuracy_mean", "accuracy_std", 0.0, 0.32, "2563EB"),
        (850, 160, 690, 400, "Final average forgetting", "forgetting_mean", "forgetting_std", 0.0, 0.70, "D97706"),
    ]
    buffers = [int(row["buffer_size"]) for row in rows]
    for x, y, width, height, title, mean_key, std_key, y_min, y_max, color in panels:
        draw.text((x, y - 5), title, font=font(28, True), fill=pil(BLACK))
        plot_top = y + 58
        plot_bottom = y + height
        plot_left = x + 82
        plot_right = x + width - 28
        draw.line((plot_left, plot_top, plot_left, plot_bottom), fill=pil(MUTED), width=2)
        draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill=pil(MUTED), width=2)
        for tick in range(5):
            value = y_min + (y_max - y_min) * tick / 4
            yy = plot_bottom - (plot_bottom - plot_top) * tick / 4
            draw.line((plot_left, yy, plot_right, yy), fill=pil("E5E7EB"), width=2)
            label = f"{value * 100:.0f}%"
            bbox = draw.textbbox((0, 0), label, font=font(19))
            draw.text((plot_left - 14 - (bbox[2] - bbox[0]), yy - 11), label, font=font(19), fill=pil(MUTED))
        point_xs = [plot_left + (plot_right - plot_left) * index / (len(rows) - 1) for index in range(len(rows))]
        points = []
        for row, point_x in zip(rows, point_xs):
            mean = float(row[mean_key])
            std = float(row[std_key])
            point_y = plot_bottom - (mean - y_min) / (y_max - y_min) * (plot_bottom - plot_top)
            top_y = plot_bottom - (mean + std - y_min) / (y_max - y_min) * (plot_bottom - plot_top)
            bottom_y = plot_bottom - (mean - std - y_min) / (y_max - y_min) * (plot_bottom - plot_top)
            points.append((point_x, point_y))
            draw.line((point_x, top_y, point_x, bottom_y), fill=pil(color), width=4)
            draw.line((point_x - 10, top_y, point_x + 10, top_y), fill=pil(color), width=4)
            draw.line((point_x - 10, bottom_y, point_x + 10, bottom_y), fill=pil(color), width=4)
        draw.line(points, fill=pil(color), width=6, joint="curve")
        for row, point_x, (_, point_y), buffer_size in zip(rows, point_xs, points, buffers):
            mean = float(row[mean_key])
            draw.ellipse((point_x - 9, point_y - 9, point_x + 9, point_y + 9), fill=pil(color))
            value_label = f"{mean * 100:.1f}%"
            bbox = draw.textbbox((0, 0), value_label, font=font(19, True))
            draw.text((point_x - (bbox[2] - bbox[0]) / 2, point_y - 42), value_label, font=font(19, True), fill=pil(color))
            buffer_label = f"{buffer_size:,}"
            bbox = draw.textbbox((0, 0), buffer_label, font=font(20))
            draw.text((point_x - (bbox[2] - bbox[0]) / 2, plot_bottom + 18), buffer_label, font=font(20), fill=pil(BLACK))
        axis_label = "Replay buffer capacity"
        bbox = draw.textbbox((0, 0), axis_label, font=font(20))
        draw.text((plot_left + (plot_right - plot_left - (bbox[2] - bbox[0])) / 2, plot_bottom + 57), axis_label, font=font(20), fill=pil(MUTED))
    image.save(CHART_PNG, dpi=(220, 220))


def add_figure(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(str(CHART_PNG), width=Inches(6.32))
    doc_pr = run._element.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", "Line charts showing final average accuracy increasing and forgetting decreasing as replay buffer capacity grows from 1,000 to 5,000, with sample-standard-deviation error bars across three seeds.")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = False
    run = caption.add_run("Figure 1. Replay-capacity results across seeds 0, 1, and 42.")
    set_run_font(run, size=8.8, italic=True, color=MUTED)


def build_report():
    rows = load_summary()
    make_chart(rows)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "A Reproducible Study of Catastrophic Forgetting in Class-Incremental Learning on CIFAR-100"
    doc.core_properties.subject = "Continual-learning research summary"
    doc.core_properties.author = "Independent Research Project"
    doc.core_properties.keywords = "continual learning, CIFAR-100, catastrophic forgetting, experience replay, EWC"

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("RESEARCH SUMMARY")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("A Reproducible Study of Catastrophic Forgetting in Class-Incremental Learning on CIFAR-100")
    set_run_font(run, size=23, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Baselines, Replay-Buffer Ablation, and Multi-Seed Validation")
    set_run_font(run, size=13.5, color=MUTED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(10)
    for label, value in (
        ("Purpose", "Research discussion at NTU CCDS"),
        ("Setting", "Single-head class-incremental learning"),
        ("Status", "Reproducible baseline and multi-seed ablation complete"),
    ):
        label_run = meta.add_run(f"{label}: ")
        set_run_font(label_run, size=9.5, bold=True, color=BLACK)
        value_run = meta.add_run(f"{value}\n")
        set_run_font(value_run, size=9.5, color=MUTED)

    add_callout(
        doc,
        "KEY RESULT",
        "Replay is the strongest tested baseline, and increasing memory from 1,000 to 5,000 examples raises three-seed mean final accuracy from 12.14% to 26.79% while reducing forgetting from 60.56% to 42.63%.",
    )

    add_heading(doc, "Abstract", 1)
    add_body(
        doc,
        "Continual-learning systems must acquire new knowledge without erasing what they learned earlier. This project studies catastrophic forgetting in a strict single-head class-incremental setting: CIFAR-100 is divided into ten sequential tasks, while a ResNet-18 always predicts over all 100 classes without receiving task identity at test time. I implemented and compared Naive Fine-Tuning, Experience Replay, and Online Elastic Weight Consolidation (EWC), retaining complete accuracy matrices and using deterministic CUDA execution for fair comparisons. At seed 42, Replay with a 2,000-example memory reached 17.96% final average accuracy, compared with 6.64% for Naive and 6.37% for Online EWC. Across seeds 0, 1, and 42, mean final accuracy increased from 12.14% at buffer 1,000 to 26.79% at buffer 5,000. The results identify replay capacity as a robust determinant of performance in this configuration while exposing persistent forgetting and an unresolved accuracy-memory trade-off.",
    )

    add_heading(doc, "1. Research question and contribution", 1)
    add_body(
        doc,
        "Research question: How does replay-buffer capacity affect catastrophic forgetting in single-head class-incremental image classification?",
        bold_prefix="Research question:",
    )
    for item in (
        "A deterministic ten-task CIFAR-100 benchmark with full post-task evaluation.",
        "Comparable Naive, reservoir-sampling Replay, and Online EWC implementations.",
        "A controlled buffer-size ablation with three-seed uncertainty estimates.",
        "Explicit retention of failed or invalid experiments instead of selective reporting.",
    ):
        add_bullet(doc, item)

    add_page_break(doc)
    add_heading(doc, "2. Experimental setup", 1)
    add_table(
        doc,
        ("Component", "Configuration"),
        (
            ("Dataset and stream", "CIFAR-100; ten tasks; ten randomly ordered classes per task"),
            ("Evaluation", "One shared 100-class head; no task identity at test time"),
            ("Backbone", "ResNet-18 adapted for 32x32 images"),
            ("Optimization", "10 epochs/task; SGD; LR 0.1; momentum 0.9; weight decay 5e-4"),
            ("Batch policy", "Current-task batch 128; replay batch 64"),
            ("Replay", "Fixed-capacity reservoir sampling"),
            ("EWC", "Online diagonal Fisher; lambda 10; decay 0.9"),
            ("Reproducibility", "Fixed class order per seed; deterministic CUDA algorithms"),
        ),
        (2300, 7060),
    )
    add_body(
        doc,
        "Let A[k,j] denote accuracy on task j after training task k. Average Accuracy is the mean of A[k,j] over learned tasks. Average Forgetting is the mean, over previous tasks, of the best historical accuracy minus current accuracy. Every run stores the accuracy matrix, complete configuration, training history, checkpoint, and summary metrics.",
    )

    add_heading(doc, "3. Results", 1)
    add_heading(doc, "3.1 Baseline comparison", 2)
    add_table(
        doc,
        ("Method, seed 42", "Final accuracy", "Final forgetting", "Time"),
        (
            ("Naive Fine-Tuning", "6.64%", "59.32%", "11.3 min"),
            ("Replay, buffer 2,000", "17.96%", "54.90%", "14.0 min"),
            ("Online EWC, lambda 10", "6.37%", "57.46%", "12.1 min"),
        ),
        (3480, 1960, 1960, 1960),
        numeric_columns=(1, 2, 3),
    )
    add_body(
        doc,
        "Naive Fine-Tuning exhibits near-complete forgetting: Task 1 falls from 54.4% immediately after it is learned to 0% after Task 2. Replay retains 40.9% on Task 1 after Task 2 and finishes at roughly 2.7 times the final average accuracy of Naive. Online EWC is stable at lambda 10 but provides no meaningful improvement in this configuration; this is a setting-specific result, not a general claim that EWC is ineffective.",
    )

    add_page_break(doc)
    add_heading(doc, "3.2 Replay capacity and multi-seed validation", 2)
    add_figure(doc)
    add_table(
        doc,
        ("Buffer", "Final accuracy, mean ± sample std", "Final forgetting, mean ± sample std"),
        (
            ("1,000", "12.14% ± 1.40%", "60.56% ± 2.47%"),
            ("2,000", "16.85% ± 1.18%", "54.36% ± 0.81%"),
            ("5,000", "26.79% ± 1.56%", "42.63% ± 1.51%"),
        ),
        (1500, 3930, 3930),
        numeric_columns=(0, 1, 2),
    )
    add_body(
        doc,
        "The final accuracies for buffers 1,000 / 2,000 / 5,000 are 11.39 / 15.62 / 25.05% for seed 0, 13.75 / 16.98 / 28.07% for seed 1, and 11.28 / 17.96 / 27.26% for seed 42. Every seed preserves the same capacity ordering. This is stronger than a single best run, but three seeds remain insufficient for precise population-level claims or formal significance testing.",
    )
    add_callout(
        doc,
        "INTERPRETATION",
        "Larger memory consistently improves both retention and final accuracy, with no visible saturation by 5,000 examples. The gain therefore remains useful but must be weighed against storage and replay cost.",
    )

    add_heading(doc, "4. Failure analysis", 1)
    add_body(
        doc,
        "The first nominal comparison was rejected because GPU nondeterminism caused Naive and Replay to diverge on Task 1 even though the replay buffer was empty. Deterministic algorithms were then enabled, after which the trajectories matched as expected. A separate EWC run with lambda 1,000 diverged at Task 5 and produced non-finite loss; the run was rejected, and training now fails immediately on non-finite values.",
    )

    add_page_break(doc)
    add_heading(doc, "5. Limitations and next question", 1)
    for item in (
        "The full Naive/Replay/EWC method comparison is still seed-42 only.",
        "The buffer stores already augmented float32 tensors, so replayed images do not receive fresh stochastic augmentation.",
        "Hyperparameters were not selected with a separate validation stream.",
        "Three seeds estimate variability but do not establish statistical significance.",
    ):
        add_bullet(doc, item)

    add_body(
        doc,
        "The next method-level experiment fixes the buffer at 2,000 examples and compares the current uniform replay baseline with one curriculum- or importance-aware sampling rule. This isolates whether which examples are replayed changes retention and positive transfer beyond merely storing more data, motivated by work on replay curricula and dynamic continual data selection. The comparison will keep the seeds, replay batch size, and total training budget paired. DER++ remains a useful secondary baseline, and a complementary systems experiment will compare float32 tensor memory with compact uint8 source-image storage and fresh replay-time augmentation.",
    )

    add_heading(doc, "6. Reproducibility", 1)
    for item in (
        "Entrypoints: train.py, plot_results.py, and run_replay_ablation.py.",
        "Artifacts: results.json, accuracy_matrix.csv, checkpoint.pt, and SVG plots for every formal run.",
        "Validation: eight automated tests covering task construction, metrics, Replay, EWC, ablation matching, statistics, and SVG output.",
        "Environment: PyTorch/CUDA runs in the validated ece488_clip Conda environment; deterministic mode is enabled by default.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "References", 1)
    references = (
        "A. Krizhevsky. Learning Multiple Layers of Features from Tiny Images. Technical report, 2009.",
        "K. He, X. Zhang, S. Ren, and J. Sun. Deep Residual Learning for Image Recognition. CVPR, 2016.",
        "J. Kirkpatrick et al. Overcoming Catastrophic Forgetting in Neural Networks. PNAS, 2017.",
        "J. S. Vitter. Random Sampling with a Reservoir. ACM Transactions on Mathematical Software, 1985.",
        "A. Chaudhry et al. Riemannian Walk for Incremental Learning: Understanding Forgetting and Intransigence. ECCV, 2018.",
        "R. J. Tee and M. Zhang. Integrating Curricula with Replays: Its Effects on Continual Learning. AAAI Symposium Series, 2023.",
        "A. Maharana et al. Adapt-infinity: Scalable Continual Multimodal Instruction Tuning via Dynamic Data Selection. ICLR, 2025.",
    )
    for reference in references:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(reference)
        set_run_font(run, size=9.2)

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)
    print(CHART_PNG)


if __name__ == "__main__":
    build_report()
